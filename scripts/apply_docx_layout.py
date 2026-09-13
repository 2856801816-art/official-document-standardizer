#!/usr/bin/env python3
"""Apply the supplied official-document layout to DOCX without rewriting text."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

BODY = "仿宋_GB2312"


def set_run_font(run, east_asia: str, size: float, bold: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)


def paragraph_role(index: int, text: str) -> tuple[str, float, bool, int]:
    stripped = text.strip()
    if index == 0 and stripped:
        return "方正小标宋简体", 22, False, WD_ALIGN_PARAGRAPH.CENTER
    if re.match(r"^[一二三四五六七八九十]+、", stripped):
        return "黑体", 16, False, WD_ALIGN_PARAGRAPH.JUSTIFY
    if re.match(r"^（[一二三四五六七八九十]+）", stripped):
        return "楷体_GB2312", 16, True, WD_ALIGN_PARAGRAPH.JUSTIFY
    if re.match(r"^\d+[．.]", stripped):
        return BODY, 16, True, WD_ALIGN_PARAGRAPH.JUSTIFY
    if re.match(r"^（\d+）", stripped):
        return BODY, 16, False, WD_ALIGN_PARAGRAPH.JUSTIFY
    return BODY, 16, False, WD_ALIGN_PARAGRAPH.JUSTIFY


def style_document(doc: Document) -> None:
    for section in doc.sections:
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.top_margin = section.bottom_margin = Mm(25.4)
        section.left_margin = section.right_margin = Mm(31.7)
        section.header_distance = section.footer_distance = Mm(12.7)
    for index, paragraph in enumerate(doc.paragraphs):
        font, size, bold, alignment = paragraph_role(index, paragraph.text)
        paragraph.alignment = alignment
        fmt = paragraph.paragraph_format
        fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        fmt.line_spacing = Pt(30)
        fmt.space_before = fmt.space_after = Pt(0)
        fmt.first_line_indent = Pt(0 if index == 0 else 32) if paragraph.text.strip() else None
        for run in paragraph.runs:
            set_run_font(run, font, size, bold)


def ensure_title(doc: Document, title: str) -> None:
    content = [paragraph for paragraph in doc.paragraphs if paragraph.text.strip()]
    has_title = bool(content and (content[0].style.name == "Title" or (content[0].alignment == WD_ALIGN_PARAGRAPH.CENTER and len(content[0].text.strip()) <= 50)))
    if has_title:
        return
    title_paragraph = doc.add_paragraph()
    title_paragraph._p.getparent().remove(title_paragraph._p)
    first = content[0]._p if content else doc._element.body.sectPr
    first.addprevious(title_paragraph._p)
    if "Title" in [style.name for style in doc.styles]:
        title_paragraph.style = doc.styles["Title"]
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_paragraph.paragraph_format.first_line_indent = Pt(0)
    title_paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    title_paragraph.paragraph_format.line_spacing = Pt(30)
    run = title_paragraph.add_run(title)
    set_run_font(run, "方正小标宋简体", 22, False)


def add_top_bottom_picture(paragraph, image_path: Path, width_mm: float = 146.0) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = Pt(6)
    shape = paragraph.add_run().add_picture(str(image_path), width=Mm(width_mm))
    inline = shape._inline
    anchor = OxmlElement("wp:anchor")
    for name, value in {"distT": "0", "distB": "0", "distL": "0", "distR": "0", "simplePos": "0", "relativeHeight": "251658240", "behindDoc": "0", "locked": "0", "layoutInCell": "1", "allowOverlap": "0"}.items():
        anchor.set(name, value)
    simple = OxmlElement("wp:simplePos")
    simple.set("x", "0")
    simple.set("y", "0")
    anchor.append(simple)
    pos_h = OxmlElement("wp:positionH")
    pos_h.set("relativeFrom", "margin")
    align_h = OxmlElement("wp:align")
    align_h.text = "center"
    pos_h.append(align_h)
    anchor.append(pos_h)
    pos_v = OxmlElement("wp:positionV")
    pos_v.set("relativeFrom", "paragraph")
    offset_v = OxmlElement("wp:posOffset")
    offset_v.text = "0"
    pos_v.append(offset_v)
    anchor.append(pos_v)
    for tag in ("wp:extent", "wp:effectExtent"):
        child = inline.find(qn(tag))
        if child is not None:
            anchor.append(child)
    anchor.append(OxmlElement("wp:wrapTopAndBottom"))
    for tag in ("wp:docPr", "wp:cNvGraphicFramePr", "a:graphic"):
        child = inline.find(qn(tag))
        if child is not None:
            anchor.append(child)
    inline.getparent().replace(inline, anchor)


def insert_after(paragraph, image_path: Path) -> None:
    new_paragraph = paragraph._parent.add_paragraph()
    new_paragraph._p.getparent().remove(new_paragraph._p)
    paragraph._p.addnext(new_paragraph._p)
    add_top_bottom_picture(new_paragraph, image_path)


def apply_plan(doc: Document, plan_path: Path) -> None:
    for item in json.loads(plan_path.read_text(encoding="utf-8")):
        needle = item["after_contains"]
        matches = [paragraph for paragraph in doc.paragraphs if needle in paragraph.text]
        if len(matches) != 1:
            raise ValueError(f"after_contains must match exactly one paragraph: {needle!r}; got {len(matches)}")
        image_path = (plan_path.parent / item["path"]).resolve()
        if not image_path.is_file():
            raise FileNotFoundError(image_path)
        insert_after(matches[0], image_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--image-plan", type=Path)
    parser.add_argument("--source-filename", help="Original filename used to supply a missing title")
    args = parser.parse_args()
    if args.input.suffix.lower() != ".docx" or args.output.suffix.lower() != ".docx":
        parser.error("This helper accepts DOCX only; preserve DOC by using a temporary DOCX and save back through Word.")
    doc = Document(args.input)
    source_name = Path(args.source_filename).stem if args.source_filename else args.output.stem
    ensure_title(doc, source_name)
    style_document(doc)
    if args.image_plan:
        apply_plan(doc, args.image_plan)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)


if __name__ == "__main__":
    main()
